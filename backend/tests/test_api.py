import io
import time

from openpyxl import load_workbook

from .conftest import login, wait_run


# ---------------------------------------------------------------- security
def test_health_and_ready(client):
    assert client.get("/api/health").json()["status"] == "ok"
    assert client.get("/api/ready").json()["database"] == "ok"


def test_requires_auth(client):
    assert client.get("/api/v1/decisions").status_code == 401
    assert client.get("/api/v1/decisions", headers={"Authorization": "Bearer junk"}).status_code == 401


def test_bad_password_rejected(client):
    r = client.post("/api/v1/auth/login", json={"email": "ceo@futuremirror.demo", "password": "wrong"})
    assert r.status_code == 401


def test_rbac_viewer_cannot_create_or_run(client):
    v = login(client, "viewer@futuremirror.demo")
    r = client.post("/api/v1/decisions", json={"title": "x" * 5, "question": "Reduce cost 10% in 6 months"}, headers=v)
    assert r.status_code == 403
    assert client.get("/api/v1/admin/users", headers=v).status_code == 403


def test_security_headers(client, ceo):
    r = client.get("/api/v1/auth/me", headers=ceo)
    assert r.headers["x-content-type-options"] == "nosniff" and r.headers["x-frame-options"] == "DENY"


def test_tenant_isolation(client, analyzed_decision):
    from app.db import SessionLocal
    from app.models import Role, User
    from app.security import hash_password
    from app.seed import seed_tenant
    with SessionLocal() as db:
        t = seed_tenant(db, slug="other-co", name="Other Co", with_users=False)
        role = db.query(Role).filter_by(tenant_id=t.id, key="admin").one()
        u = User(tenant_id=t.id, email="admin@other.co", full_name="Other Admin", password_hash=hash_password("FutureMirror!2026"))
        u.roles = [role]
        db.add(u)
        db.commit()
    other = login(client, "admin@other.co", tenant="other-co")
    did = analyzed_decision["decision"]["id"]
    assert client.get(f"/api/v1/decisions/{did}", headers=other).status_code == 404
    assert client.get(f"/api/v1/decisions/{did}/dossier", headers=other).status_code == 404
    assert all(d["id"] != did for d in client.get("/api/v1/decisions", headers=other).json())


# ---------------------------------------------------------------- decision pipeline
def test_pipeline_produces_full_dossier(client, analyst, analyzed_decision):
    did = analyzed_decision["decision"]["id"]
    ds = client.get(f"/api/v1/decisions/{did}/dossier", headers=analyst).json()
    assert ds["dataset"]["is_synthetic"] is True
    assert ds["run"]["mode"] == "rule_based"
    agents = {f["agent_key"] for f in ds["findings"]}
    assert {"strategy", "cfo", "coo", "hr", "delivery", "legal", "risk", "governance"} <= agents
    assert len(ds["objections"]) >= 3 and all(o["raised_by"] == "redteam" for o in ds["objections"])
    assert any(c["rule_code"] == "COL-01" for c in ds["collisions"])
    assert set(ds["simulation"]["results"]) == {"BASELINE", "A", "B", "C"}
    assert ds["simulation"]["model"]["version"]
    assert len(ds["executions"]) >= 15
    rec = ds["run"]["result"]["recommendation"]
    assert rec["per_future"]["A"]["governance_outcome"] == "denied"
    assert ds["run"]["result"]["governance_probe"]["outcome"] == "denied"
    assert ds["decision"]["status"] == "analyzed"


def test_events_cover_all_stages_and_are_real(client, analyst, analyzed_decision):
    run = client.get(f"/api/v1/runs/{analyzed_decision['run_id']}", headers=analyst).json()
    started = [e["stage"] for e in run["events"] if e["event_type"] == "stage_started"]
    for st in ("intake", "data_validation", "agent_assignment", "independent_analysis", "scenario_generation",
               "quantitative_simulation", "agent_court", "risk_assessment", "governance_review", "executive_recommendation"):
        assert st in started
    assert any(e["event_type"] == "run_completed" for e in run["events"])
    assert run["compute_ms"] > 0


def test_sse_stream_replays_events(client, analyst, analyzed_decision):
    with client.stream("GET", f"/api/v1/runs/{analyzed_decision['run_id']}/events", headers=analyst) as r:
        body = ""
        for chunk in r.iter_text():
            body += chunk
            if "event: end" in body:
                break
    assert "event: workflow" in body and "stage_started" in body and "event: end" in body


def test_concurrent_run_rejected_and_cancel(client, analyst):
    d = client.post("/api/v1/decisions", json={"title": "Cancel me", "question": "Reduce costs by 10% in 6 months"},
                    headers=analyst).json()
    r = client.post(f"/api/v1/decisions/{d['id']}/runs", json={"pacing_ms": 1500}, headers=analyst).json()
    assert client.post(f"/api/v1/decisions/{d['id']}/runs", json={}, headers=analyst).status_code == 409
    ceo = login(client, "ceo@futuremirror.demo")
    assert client.post(f"/api/v1/runs/{r['run_id']}/cancel", headers=ceo).status_code == 200
    res = wait_run(client, analyst, r["run_id"])
    assert res["status"] == "cancelled"


def test_agent_retry_and_failure_handling(client, analyst):
    from app.db import SessionLocal
    from app.models import AgentExecution, Decision, DecisionRun
    from app.services.pipeline import execute_run
    d = client.post("/api/v1/decisions", json={"title": "Retry test", "question": "Reduce costs by 10% in 6 months"},
                    headers=analyst).json()
    with SessionLocal() as db:
        run = DecisionRun(decision_id=d["id"], status="queued", pacing_ms=0)
        db.add(run)
        db.commit()
        rid = run.id
    execute_run(rid, fail_injection={"hr": 1, "legal": 5})
    r = client.get(f"/api/v1/runs/{rid}", headers=analyst).json()
    types = [(e["event_type"], e["agent_key"]) for e in r["events"]]
    assert ("agent_retry", "hr") in types
    assert ("agent_failed", "legal") in types
    assert r["status"] == "completed"  # pipeline degrades gracefully
    with SessionLocal() as db:
        ex = db.query(AgentExecution).filter_by(run_id=rid, agent_key="legal", phase="independent_analysis").one()
        assert ex.status == "failed" and ex.attempts == 3
        assert db.get(Decision, d["id"]).status == "analyzed"


def test_emergency_stop_blocks_runs(client, analyst):
    admin = login(client, "admin@futuremirror.demo")
    assert client.post("/api/v1/admin/emergency-stop", json={"active": True, "reason": "test"}, headers=admin).status_code == 200
    d = client.post("/api/v1/decisions", json={"title": "Stopped", "question": "Reduce costs by 10% in 6 months"},
                    headers=analyst).json()
    assert client.post(f"/api/v1/decisions/{d['id']}/runs", json={}, headers=analyst).status_code == 423
    client.post("/api/v1/admin/emergency-stop", json={"active": False}, headers=admin)
    assert client.get("/api/v1/platform/status", headers=admin).json()["emergency_stop"] is False


# ---------------------------------------------------------------- simulation API
def test_what_if_recalculates(client, analyst):
    a = client.post("/api/v1/simulation/run", json={"iterations": 300, "include_sensitivity": False,
                                                    "include_des": False}, headers=analyst).json()
    b = client.post("/api/v1/simulation/run", json={"iterations": 300, "include_sensitivity": False, "include_des": False,
                                                    "controls": {"automation_investment": 12}}, headers=analyst).json()
    assert a["meta"]["input_fingerprint"] != b["meta"]["input_fingerprint"]
    assert (b["futures"]["B"]["summary"]["one_off_implementation_cost"]["expected"]
            > a["futures"]["B"]["summary"]["one_off_implementation_cost"]["expected"])


def test_invalid_controls_rejected(client, analyst):
    r = client.post("/api/v1/simulation/run", json={"controls": {"cost_reduction_target_pct": 500}}, headers=analyst)
    assert r.status_code == 422


def test_custom_lab_strategy(client, analyst):
    r = client.post("/api/v1/simulation/custom", json={"levers": {"software_consolidation": 1, "infra_optimization": 1},
                                                       "iterations": 300}, headers=analyst)
    assert r.status_code == 200 and r.json()["summary"]["max_delay_increase_pct"]["p90"] == 0


# ---------------------------------------------------------------- governance & approvals
def test_locked_policy_cannot_be_changed(client):
    admin = login(client, "admin@futuremirror.demo")
    pols = client.get("/api/v1/governance/policies", headers=admin).json()
    locked = next(p for p in pols if p["code"] == "GOV-001")
    assert client.patch(f"/api/v1/governance/policies/{locked['id']}", json={"enabled": False}, headers=admin).status_code == 403


def test_blocked_strategy_cannot_enter_approval(client, analyst, analyzed_decision):
    did = analyzed_decision["decision"]["id"]
    r = client.post(f"/api/v1/decisions/{did}/approval-requests",
                    json={"scenario_key": "A", "acknowledged": ["GOV-011"]}, headers=analyst)
    assert r.status_code == 201 and r.json()["status"] == "blocked"


def test_acknowledgement_required(client, analyst, analyzed_decision):
    did = analyzed_decision["decision"]["id"]
    r = client.post(f"/api/v1/decisions/{did}/approval-requests", json={"scenario_key": "B"}, headers=analyst)
    assert r.status_code == 422


def test_approval_workflow_with_separation_of_duties(client, analyst, analyzed_decision):
    did = analyzed_decision["decision"]["id"]
    r = client.post(f"/api/v1/decisions/{did}/approval-requests",
                    json={"scenario_key": "B", "acknowledged": ["GOV-011"], "justification": "Basis for revised target"},
                    headers=analyst)
    assert r.status_code == 201, r.text
    req = r.json()
    assert req["status"] == "pending" and set(req["required_roles"]) >= {"cfo", "legal"}
    rid = req["id"]
    viewer = login(client, "viewer@futuremirror.demo")
    assert client.post(f"/api/v1/approvals/{rid}/decide", json={"verdict": "approve", "role_key": "cfo"},
                       headers=viewer).status_code == 403
    cfo = login(client, "cfo@futuremirror.demo")
    assert client.post(f"/api/v1/approvals/{rid}/decide", json={"verdict": "approve", "role_key": "legal"},
                       headers=cfo).status_code == 403  # does not hold legal role
    users = {"cfo": cfo, "legal": login(client, "legal@futuremirror.demo"), "ceo": login(client, "ceo@futuremirror.demo"),
             "coo": login(client, "coo@futuremirror.demo"), "hr": login(client, "hr@futuremirror.demo")}
    for role in req["required_roles"]:
        r = client.post(f"/api/v1/approvals/{rid}/decide", json={"verdict": "approve", "role_key": role,
                                                                "comment": "ok"}, headers=users[role])
        assert r.status_code == 200, r.text
    final = r.json()
    assert final["status"] == "approved" and final["pending_roles"] == []
    assert client.get(f"/api/v1/decisions/{did}", headers=analyst).json()["status"] == "approved"
    # double approval rejected
    assert client.post(f"/api/v1/approvals/{rid}/decide", json={"verdict": "approve", "role_key": "cfo"},
                       headers=cfo).status_code == 409


def test_requester_cannot_approve_own_request(client):
    cfo = login(client, "cfo@futuremirror.demo")
    q = "Reduce operating costs by 10% within 6 months without layoffs"
    d = client.post("/api/v1/decisions", json={"title": "SoD check", "question": q}, headers=cfo).json()
    run = client.post(f"/api/v1/decisions/{d['id']}/runs", json={"pacing_ms": 0}, headers=cfo).json()
    wait_run(client, cfo, run["run_id"])
    r = client.post(f"/api/v1/decisions/{d['id']}/approval-requests",
                    json={"scenario_key": "B", "acknowledged": ["GOV-011"]}, headers=cfo).json()
    resp = client.post(f"/api/v1/approvals/{r['id']}/decide", json={"verdict": "approve", "role_key": "cfo"}, headers=cfo)
    assert resp.status_code == 403 and "GOV-013" in resp.json()["detail"]


# ---------------------------------------------------------------- automation, feedback, reports
def test_automation_dry_run_idempotent_and_live_blocked(client):
    cfo = login(client, "cfo@futuremirror.demo")
    wfs = client.get("/api/v1/automation/workflows", headers=cfo).json()["workflows"]
    kpi = next(w for w in wfs if w["key"] == "kpi_monitor")
    r1 = client.post(f"/api/v1/automation/workflows/{kpi['id']}/trigger", json={"dry_run": True, "payload": {"x": 1}},
                     headers=cfo).json()
    r2 = client.post(f"/api/v1/automation/workflows/{kpi['id']}/trigger", json={"dry_run": True, "payload": {"x": 1}},
                     headers=cfo).json()
    assert r1["status"] == "dry_run_ok" and r1["id"] == r2["id"]
    danger = next(w for w in wfs if w["consequential"])
    r3 = client.post(f"/api/v1/automation/workflows/{danger['id']}/trigger", json={"dry_run": False}, headers=cfo).json()
    assert r3["status"] == "blocked" and "GOV-014" in r3["result"]["policies"]


def test_reality_feedback_loop_and_model_governance(client, analyst, analyzed_decision):
    did = analyzed_decision["decision"]["id"]
    out = client.post(f"/api/v1/decisions/{did}/outcomes/synthetic-demo", headers=analyst).json()
    assert len(out["comparisons"]) >= 10 and all(a["is_synthetic"] for a in out["actuals"])
    cal = out["calibration"]
    assert cal["metrics"]["delay_pct"]["above_p90"] >= 1
    assert any(d["parameter"] == "skill_match" for d in cal["diagnostics"])
    manual = client.post(f"/api/v1/decisions/{did}/outcomes", json=[{"month": 4, "metric": "csat_pct", "value": 85.5}],
                         headers=analyst).json()
    assert any(c["month"] == 4 for c in manual["comparisons"])
    prop = client.post(f"/api/v1/decisions/{did}/calibration-proposal", headers=analyst).json()
    assert prop["created"] and prop["candidate"]["status"] == "candidate"
    model = client.get("/api/v1/simulation/model", headers=analyst).json()
    assert model["production"]["status"] == "production" and model["production"]["id"] != prop["candidate"]["id"]
    assert client.post(f"/api/v1/models/{prop['candidate']['id']}/promote", json={"justification": "Validated on 3 months"},
                       headers=analyst).status_code == 403
    admin = login(client, "admin@futuremirror.demo")
    pr = client.post(f"/api/v1/models/{prop['candidate']['id']}/promote", json={"justification": "Validated on 3 months"},
                     headers=admin)
    assert pr.status_code == 200 and pr.json()["status"] == "production"


def test_reports_pdf_and_excel_both_languages(client, analyst, analyzed_decision):
    did = analyzed_decision["decision"]["id"]
    for lang in ("en", "ar"):
        pdf = client.get(f"/api/v1/decisions/{did}/report.pdf?lang={lang}", headers=analyst)
        assert pdf.status_code == 200 and pdf.content[:4] == b"%PDF" and len(pdf.content) > 5000
        xl = client.get(f"/api/v1/decisions/{did}/report.xlsx?lang={lang}", headers=analyst)
        wb = load_workbook(io.BytesIO(xl.content))
        assert "Futures" in wb.sheetnames and "Agent findings" in wb.sheetnames


def test_dataset_versioning_and_import(client, analyst):
    ds = client.get("/api/v1/datasets", headers=analyst).json()[0]
    v = client.post(f"/api/v1/datasets/{ds['id']}/versions", json={"data": {"gross_margin_pct": 70}, "note": "fix margin"},
                    headers=analyst)
    assert v.status_code == 201 and v.json()["provenance"]["gross_margin_pct"] == "user_input"
    dq2 = next(c for c in v.json()["validation"]["checks"] if c["code"] == "DQ-002")
    assert dq2["severity"] == "ok"
    csv_body = "category,amount\nsoftware licensing,10\ncloud,8.5\nmystery,1\n"
    r = client.post(f"/api/v1/datasets/{ds['id']}/import", files={"file": ("gl.csv", csv_body, "text/csv")}, headers=analyst)
    assert r.status_code == 201, r.text
    rec = r.json()["reconciliation"]
    assert rec["unmapped_categories"] == {"mystery": 1.0} and r.json()["version"]["data"]["software_cost"] == 10
    bad = client.post(f"/api/v1/datasets/{ds['id']}/import", files={"file": ("x.csv", "category,amount\na,notnum\n", "text/csv")},
                      headers=analyst)
    assert bad.status_code == 422


def test_connectors_honest_status(client):
    admin = login(client, "admin@futuremirror.demo")
    conns = {c["kind"]: c for c in client.get("/api/v1/connectors", headers=admin).json()}
    assert conns["synthetic_erp"]["implementation"] == "working"
    sap = client.post(f"/api/v1/connectors/{conns['sap']['id']}/test", headers=admin).json()
    assert sap["ok"] is False and sap["verified_live"] is False
    odoo = client.post(f"/api/v1/connectors/{conns['odoo']['id']}/test", headers=admin).json()
    assert odoo["ok"] is False
    syn = client.post(f"/api/v1/connectors/{conns['synthetic_erp']['id']}/sync", headers=admin).json()
    assert abs(syn["reconciliation"]["difference"]) < 1e-6 and syn["verified_live"] is False
    r = client.patch(f"/api/v1/connectors/{conns['odoo']['id']}", json={"config": {"api_key": "x"}}, headers=admin)
    assert r.status_code == 422  # secrets never stored in config


def test_connector_settings_validation_and_d365(client, monkeypatch):
    admin = login(client, "admin@futuremirror.demo")
    conns = {c["kind"]: c for c in client.get("/api/v1/connectors", headers=admin).json()}
    fo = conns["d365_fo"]
    assert fo["implementation"] == "adapter" and fo["secret_env"] == ["FM_D365_FO_CLIENT_SECRET"]
    assert conns["odoo"]["secret_env"] == ["FM_ODOO_API_KEY"]
    assert {conns[k]["implementation"] for k in ("sap", "rest", "sql")} == {"adapter"}
    assert conns["sap"]["secret_env"] == ["FM_SAP_PASSWORD"] and conns["rest"]["secret_optional"] is True
    url = f"/api/v1/connectors/{fo['id']}"
    for bad in ({"environment_url": "javascript:alert(1)"}, {"date_from": "01/01/2026"},
                {"date_from": "2026-12-31", "date_to": "2026-01-01"}, {"account_mapping": {"61": "salary-ish"}},
                {"account_mapping": ["61"]}, {"legal_entity": "eg01' or 1 eq 1"}, {"client_secret": "x"},
                {"dialect": "oracle"}, {"port": "5432"}, {"strip_leading_zeros": "yes"}):
        assert client.patch(url, json={"config": bad}, headers=admin).status_code == 422, bad
    analyst = login(client, "analyst@futuremirror.demo")
    assert client.patch(url, json={"config": {"tenant_id": "t"}}, headers=analyst).status_code == 403
    cfg = {"environment_url": "https://contoso.operations.dynamics.com", "tenant_id": "tid", "client_id": "cid",
           "date_from": "2026-01-01", "date_to": "2026-12-31", "account_mapping": {"61": "personnel"}}
    saved = client.patch(url, json={"config": cfg}, headers=admin).json()
    assert saved["config"]["account_mapping"] == {"61": "personnel"} and saved["secret_configured"] is False
    r = client.post(f"{url}/test", headers=admin).json()
    assert r["ok"] is False and "FM_D365_FO_CLIENT_SECRET" in r["message"]
    monkeypatch.setenv("FM_D365_FO_CLIENT_SECRET", "s3cret")
    listed = {c["kind"]: c for c in client.get("/api/v1/connectors", headers=admin).json()}
    assert listed["d365_fo"]["secret_configured"] is True and "s3cret" not in str(listed)


def test_audit_chain_valid(client):
    admin = login(client, "admin@futuremirror.demo")
    a = client.get("/api/v1/admin/audit", headers=admin).json()
    assert a["chain"]["valid"] and a["chain"]["checked"] > 10


def test_persistence_survives_restart(client, analyst, analyzed_decision):
    """A fresh application instance on the same database sees all persisted results."""
    from fastapi.testclient import TestClient

    from app.main import create_app
    with TestClient(create_app()) as c2:
        ds = c2.get(f"/api/v1/decisions/{analyzed_decision['decision']['id']}/dossier", headers=analyst).json()
        assert ds["run"]["status"] == "completed" and len(ds["findings"]) > 20


def test_dashboard(client, ceo):
    d = client.get("/api/v1/dashboard", headers=ceo).json()
    assert "decision_counts" in d and d["total_runs"] >= 1
    time.sleep(0)
