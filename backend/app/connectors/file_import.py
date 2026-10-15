"""CSV / Excel import with mapping, validation and reconciliation.

Two accepted layouts (auto-detected by header):
  * key/value profile:  ``field,value``           e.g. ``annual_revenue,180``
  * ledger lines:       ``category,amount`` (+ optional ``period``) aggregated per category
Amounts are EGP millions unless a ``unit`` column says ``EGP`` (then divided by 1e6).
"""

from __future__ import annotations

import csv
import io
from openpyxl import load_workbook
from pydantic import ValidationError

from ..simulation.assumptions import Baseline
from .base import ConnectorError, ConnectorResult, reconcile

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 50_000


def _rows_from_bytes(filename: str, content: bytes) -> list[dict[str, str]]:
    if len(content) > MAX_BYTES:
        raise ConnectorError("File too large (max 5 MB)")
    name = filename.lower()
    if name.endswith(".csv"):
        text = content.decode("utf-8-sig", errors="strict")
        reader = csv.DictReader(io.StringIO(text))
        rows = [dict(r) for _, r in zip(range(MAX_ROWS), reader, strict=False)]
    elif name.endswith(".xlsx"):
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(it)]
        rows = []
        for _, vals in zip(range(MAX_ROWS), it, strict=False):
            rows.append({header[i]: ("" if v is None else str(v)) for i, v in enumerate(vals) if i < len(header)})
    else:
        raise ConnectorError("Unsupported file type (use .csv or .xlsx)")
    return [{(k or "").strip().lower(): (v or "").strip() for k, v in r.items()} for r in rows]


def import_profile(filename: str, content: bytes, base: Baseline | None = None) -> ConnectorResult:
    base = base or Baseline()
    rows = _rows_from_bytes(filename, content)
    if not rows:
        raise ConnectorError("File contains no data rows")
    cols = set(rows[0])
    errors: list[str] = []
    if {"field", "value"} <= cols:
        data = base.model_dump()
        unknown = []
        for i, r in enumerate(rows, start=2):
            f = r["field"]
            if f not in Baseline.model_fields:
                unknown.append(f)
                continue
            try:
                data[f] = r["value"] if f == "currency" else float(r["value"].replace(",", ""))
            except ValueError:
                errors.append(f"Row {i}: value for {f} is not numeric")
        if errors:
            raise ConnectorError("; ".join(errors[:10]))
        try:
            profile = Baseline(**data).model_dump()
        except ValidationError as exc:
            raise ConnectorError(f"Validation failed: {exc.errors()[0]['msg']} ({exc.errors()[0]['loc']})") from exc
        rec = {"layout": "key_value", "unknown_fields": unknown, "fields_updated": len(rows) - len(unknown),
               "status": "ok" if not unknown else "partial"}
        return ConnectorResult(profile=profile, reconciliation=rec, rows_read=len(rows), source="file", verified_live=False)

    if {"category", "amount"} <= cols:
        totals: dict[str, float] = {}
        for i, r in enumerate(rows, start=2):
            try:
                amt = float(r["amount"].replace(",", ""))
            except ValueError:
                errors.append(f"Row {i}: amount is not numeric")
                continue
            if r.get("unit", "").upper() == "EGP":
                amt /= 1e6
            totals[r["category"]] = totals.get(r["category"], 0.0) + amt
        if errors:
            raise ConnectorError("; ".join(errors[:10]))
        profile, rec = reconcile(totals, base)
        rec["layout"] = "ledger"
        rec["note"] = "Categories not present in the file retain their previous (synthetic) values."
        try:
            Baseline(**profile)
        except ValidationError as exc:
            raise ConnectorError(f"Validation failed: {exc.errors()[0]['msg']}") from exc
        return ConnectorResult(profile=profile, reconciliation=rec, rows_read=len(rows), source="file", verified_live=False)
    raise ConnectorError("Unrecognised layout: expected columns field,value or category,amount")
